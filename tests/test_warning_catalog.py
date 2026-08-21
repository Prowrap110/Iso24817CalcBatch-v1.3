from datetime import UTC, datetime
from io import BytesIO

from openpyxl import load_workbook
import pytest

from engine.corrosion_defects import ENTER_MANUALLY
from tests.helpers import detail_values, valid_row_values, workbook_bytes_with_rows
from warning_catalog import UnmappedWarningError, warning_codes, warning_meaning
from workbook_processor import process_workbook


WARNING_SAMPLES = (
    ('W001', 'Design temperature 91.0 degC exceeds the qualified Prowrap limit of 90.00 degC.'),
    ('W002', 'NOT REPAIRABLE PER ISO 24817 FORMULA 12: no thickness can satisfy the case.'),
    ('W003', 'Type B service life is capped at 2 years for PRW110.'),
    ('W004', 'Design temperature 81.0 degC exceeds the Type B upper service limit of 80.0 degC.'),
    ('W005', 'Type B defect at zero design pressure: Formula 12 is non-controlling.'),
    ('W006', 'Type B design assumes a circular/near-circular defect of size 15 mm.'),
    ('W007', 'Formula 12 validity exceeded: defect size 100 mm exceeds the limit.'),
    ('W008', 'B31G: d/t > 0.80: beyond B31G applicability.'),
    ('W009', 'B31G: d/t <= 0.10: metal loss is not limited as to length.'),
    ('W010', 'B31G: Safety factor < 1.25 is below the minimum.'),
    ('W011', 'B31G: SMYS > 483 MPa: falling back to Original B31G.'),
    ('W012', 'B31G: Flow stress capped at SMTS.'),
    ('W013', 'B31G Level 1: the corroded pipe alone is NOT acceptable at design pressure.'),
    ('W014', 'Internal corrosion projected at 0.10 mm/yr to end of design life.'),
    ('W015', 'Internal corrosion with corrosion rate = 0 mm/yr requires review.'),
    ('W016', 'Axial load case 1 (Formula 4 end-thrust) selected with a Type B defect.'),
    ('W017', 'Repair thickness exceeds D/12.'),
    ('W018', 'Prowrap CF cloth width 250 mm is not an approved configuration.'),
    ('W019', 'Type A / Class 3 check was not run at zero design pressure.'),
    ('W020', 'Type A / Class 3 check was not run above the qualified Prowrap temperature limit.'),
)


@pytest.mark.parametrize(('expected_code', 'message'), WARNING_SAMPLES)
def test_each_engineering_warning_shape_maps_to_one_permanent_code(expected_code, message):
    """Catches a warning condition that loses or changes its permanent reference."""
    assert warning_codes((message,)) == (expected_code,)


def test_warning_codes_preserve_first_occurrence_order_and_remove_duplicates():
    """Catches repeated row warnings that would make the compact cell noisy."""
    messages = (WARNING_SAMPLES[17][1], WARNING_SAMPLES[0][1], WARNING_SAMPLES[17][1])

    assert warning_codes(messages) == ('W018', 'W001')


def test_candidate_prefixed_b31g_warnings_keep_their_permanent_codes():
    """Catches linked-detail warning prefixes breaking B31G traceability."""
    messages = (
        'Defect ID D-01: B31G: SMYS > 483 MPa: falling back to Original B31G.',
        'Defect ID D-01: B31G Level 1: the corroded pipe alone is NOT '
        'acceptable at the design pressure.',
    )

    assert warning_codes(messages) == ('W011', 'W013')


def test_unknown_warning_is_rejected_instead_of_receiving_an_improvised_code():
    """Catches uncatalogued engineering conditions that would become untraceable."""
    with pytest.raises(UnmappedWarningError, match='Unmapped compliance warning'):
        warning_codes(('New unregistered engineering warning.',))


def test_cloth_warning_meaning_names_both_approved_widths():
    """Catches a warning register that still describes only the old 300 mm approval."""
    meaning = warning_meaning('W018')

    assert '300 mm' in meaning
    assert '500 mm' in meaning


def test_warning_register_identifies_main_and_detail_rows():
    """Catches linked warnings being recorded only against their main repair."""
    source = workbook_bytes_with_rows(
        [valid_row_values(**{
            'Pipe OD [mm]': 1016.0,
            'Nominal Wall [mm]': 12.0,
            'Pipe Yield [MPa]': 450.0,
            'Design Pressure [bar]': 104.9,
            'Defect Length [mm]': 1000.0,
            'Defect Length Basis': ENTER_MANUALLY,
            'Repair Group ID': 'R-001',
            'Remaining Wall [mm]': None,
            'Prowrap CF Cloth Width [mm]': 500.0,
            'Run Type A / Class 3 Check': 'Yes',
        })],
        detail_rows=[
            detail_values(group='R-001', defect='D-01', length=10, wall=9.652),
            detail_values(group='R-001', defect='D-02', length=35, wall=10.0),
        ],
    )

    processed = process_workbook(
        source, processed_at=datetime(2026, 8, 14, 12, 0, tzinfo=UTC),
    )
    workbook = load_workbook(BytesIO(processed.workbook_bytes), data_only=False)
    warnings = workbook['Warnings']
    warning_rows = {
        warnings.cell(row, 1).value: warnings.cell(row, 3).value
        for row in range(4, warnings.max_row + 1)
    }

    assert warning_rows['W013'] == 'Main 2; Individual Defects 2, 3'
    assert workbook['Summary']['B24'].value == '1.3.0'
    assert workbook['Summary']['B25'].value == 'da83373'
