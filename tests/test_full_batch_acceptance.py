"""End-to-end release acceptance for the isolated CalcBatch v1.3 app."""

from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
import pytest

from batch_schema import DETAIL_INPUT_HEADERS, INPUT_HEADERS, OUTPUT_HEADERS
from cost_calculation import COST_SOURCE_HEADERS, cost_formula, price_formula, total_amount_formula
from workbook_processor import process_workbook


FIXED_TIME = datetime(2026, 8, 21, 12, 0, tzinfo=UTC)
EXPECTED_SHEETS = [
    'Batch Information', 'Batch Input & Results', 'Individual Defects',
    'Cost Calculation', 'Warnings', 'Summary', 'Instructions', 'Lists',
]
EXPECTED_STRUCTURAL_OUTPUTS = (
    52.7806925498426,
    2.0,
    3,
    588.933816016055,
    300.0,
)
EXPECTED_PROCUREMENT = (
    (0, 3, 900.0, 3.878107635297385, 4.6537291623568615),
    (2, 0, 1000.0, 4.30900848366376, 5.170810180396512),
    (1, 1, 800.0, 3.447206786931009, 4.136648144317211),
    (1, 1, 800.0, 3.447206786931009, 4.136648144317211),
)


def _columns(headers):
    return {header: index for index, header in enumerate(headers, start=1)}


def _formula_cells(workbook):
    return [
        (f'{worksheet.title}!{cell.coordinate}', cell.value)
        for worksheet in workbook.worksheets
        for row in worksheet.iter_rows()
        for cell in row
        if cell.data_type == 'f'
    ]


def _main_result_signature(workbook):
    worksheet = workbook['Batch Input & Results']
    columns = _columns(INPUT_HEADERS + OUTPUT_HEADERS)
    return tuple(
        tuple(worksheet.cell(row, columns[heading]).value for heading in OUTPUT_HEADERS)
        for row in range(2, 8)
    )


def _warning_signature(workbook):
    worksheet = workbook['Warnings']
    return tuple(
        tuple(worksheet.cell(row, column).value for column in range(1, 4))
        for row in range(4, worksheet.max_row + 1)
        if worksheet.cell(row, 1).value
    )


def _summary_signature(workbook):
    worksheet = workbook['Summary']
    return tuple(
        worksheet[address].value
        for address in (
            'B3', 'B4', 'B5', 'B10', 'B13', 'B14',
            'B15', 'B16', 'B17', 'B24', 'B25',
        )
    )


def test_release_documentation_is_current_v13_only_and_names_isolated_deployment():
    """Catch release guidance that points operators to v1.2 or an existing app."""
    root = Path(__file__).resolve().parents[1]
    readme = (root / 'README.md').read_text(encoding='utf-8')
    deployment = (root / 'DEPLOYMENT.md').read_text(encoding='utf-8')

    assert readme.startswith('# PROWRAP CalcBatch v1.3\n')
    assert 'PROWRAP_CalcBatch_v1.3_Template.xlsx' in readme
    assert 'PROWRAP_CalcBatch_v1.2_Template.xlsx' not in readme
    assert '21 controlled inputs and 10 controlled outputs' in readme
    assert '300-only, 500-only, mixed, and reversed mixed' in readme
    assert 'Batch release version is `1.3.0`' in readme
    assert '`da83373d648694f50b8a974ff6071a73ceec2089`' in readme

    assert deployment.startswith(
        '# Deploy CalcBatch v1.3 as a new isolated public application\n'
    )
    assert '`Prowrap110/Iso24817CalcBatch-v1.3`' in deployment
    assert '`release/v1.3.0`' in deployment
    assert '`app.py`' in deployment
    assert '`iso24817calcbatch-prowrapv13`' in deployment
    assert '`https://iso24817calcbatch-prowrapv13.streamlit.app`' in deployment
    assert 'Do not deploy now' in deployment
    assert 'Do not select, reboot, reconfigure, or deploy over' in deployment


def test_mixed_width_v13_release_acceptance_workbook(tmp_path):
    """Exercise all six frozen v1.3 scenarios through production paths."""
    from scripts.create_acceptance_workbook import create_acceptance_workbook

    source_path = tmp_path / 'PROWRAP_CalcBatch_v1.3_Acceptance_Input.xlsx'
    create_acceptance_workbook(source_path)
    input_book = load_workbook(source_path, data_only=False)
    main_input = input_book['Batch Input & Results']
    main_columns = _columns(INPUT_HEADERS + OUTPUT_HEADERS)

    assert input_book.sheetnames == EXPECTED_SHEETS
    assert input_book.properties.title == 'PROWRAP CalcBatch v1.3'
    assert input_book['Batch Information']['A1'].value == 'PROWRAP CalcBatch v1.3'
    assert input_book['Instructions']['A1'].value == (
        'PROWRAP CalcBatch v1.3 — Instructions'
    )
    assert len(INPUT_HEADERS) == 21
    assert len(OUTPUT_HEADERS) == 10
    assert tuple(cell.value for cell in main_input[1]) == INPUT_HEADERS + OUTPUT_HEADERS
    assert [
        input_book['Batch Information'].cell(row, 2).value for row in (3, 4, 5)
    ] == ['Acceptance Customer', 'Acceptance Location', 'ACCEPT-V13-001']
    assert [
        tuple(main_input.cell(row, main_columns[header]).value for header in (
            'Prowrap CF Cloth Width 1 [mm]',
            'Prowrap CF Cloth Width 2 [mm]',
        ))
        for row in range(2, 8)
    ] == [
        (300, 300), (500, 500), (300, 500),
        (500, 300), (300, 500), (300, 500),
    ]
    assert [
        main_input.cell(row, main_columns['Defect Length [mm]']).value
        for row in range(2, 8)
    ] == [300] * 6
    assert main_input.cell(6, main_columns['Pipe OD [mm]']).value is None
    assert main_input.cell(7, main_columns['Mechanism']).value == 'Leak'
    assert main_input.cell(7, main_columns['Design Pressure [bar]']).value == 150
    assert _formula_cells(input_book) == []

    assert (
        main_input.tables['BatchRows'].ref,
        main_input.tables['BatchRows'].autoFilter.ref,
    ) == ('A1:AE151', 'A1:AE151')
    detail_input = input_book['Individual Defects']
    assert (
        detail_input.tables['IndividualDefects'].ref,
        detail_input.tables['IndividualDefects'].autoFilter.ref,
    ) == ('A1:X151', 'A1:X151')
    assert tuple(
        cell.value for cell in detail_input[1][:len(DETAIL_INPUT_HEADERS)]
    ) == DETAIL_INPUT_HEADERS
    validations = {
        item.formula1: str(item.sqref)
        for item in main_input.data_validations.dataValidation
    }
    assert validations['=ClothWidth1Choices'].endswith('151')
    assert validations['=ClothWidth2Choices'].endswith('151')

    processed = process_workbook(
        source_path.read_bytes(),
        processed_at=FIXED_TIME,
        source_name=source_path.name,
    )
    result_book = load_workbook(BytesIO(processed.workbook_bytes), data_only=False)
    main = result_book['Batch Input & Results']
    detail = result_book['Individual Defects']
    cost = result_book['Cost Calculation']

    assert processed.status_counts == {
        'OK': 4, 'INPUT ERROR': 1, 'NOT REPAIRABLE': 1,
    }
    assert _summary_signature(result_book) == (
        'Acceptance Customer', 'Acceptance Location', 'ACCEPT-V13-001',
        6, 4, 0, 1, 1, 0, '1.3.0', 'da83373',
    )
    assert result_book['Lists'].sheet_state == 'hidden'
    assert all(result_book[name].protection.sheet for name in (
        'Batch Input & Results', 'Individual Defects', 'Cost Calculation',
        'Warnings', 'Summary',
    ))
    assert (
        main.protection.autoFilter,
        main.protection.selectLockedCells,
        main.protection.selectUnlockedCells,
    ) == (False, False, False)
    assert (
        detail.protection.autoFilter,
        detail.protection.selectLockedCells,
        detail.protection.selectUnlockedCells,
    ) == (False, False, False)
    assert all(
        not main.cell(2, column).protection.locked
        for column in range(1, len(INPUT_HEADERS) + 1)
    )
    assert all(
        main.cell(2, column).protection.locked
        for column in range(
            len(INPUT_HEADERS) + 1,
            len(INPUT_HEADERS + OUTPUT_HEADERS) + 1,
        )
    )
    assert all(
        not detail.cell(2, column).protection.locked
        for column in range(1, len(DETAIL_INPUT_HEADERS) + 1)
    )
    assert detail.cell(2, len(DETAIL_INPUT_HEADERS) + 1).protection.locked

    structural_headers = (
        'Wall Loss [%]', 'Required Structural Thickness [mm]',
        'Installed Plies', 'Total Repair Length [mm]',
        'Repair Zone Length [mm]',
    )
    for row in range(2, 6):
        assert tuple(
            main.cell(row, main_columns[header]).value
            for header in structural_headers
        ) == pytest.approx(EXPECTED_STRUCTURAL_OUTPUTS)

    for row, expected in zip(range(2, 6), EXPECTED_PROCUREMENT, strict=True):
        actual = tuple(
            main.cell(row, main_columns[header]).value
            for header in (
                '500 mm Cloth Band Count', '300 mm Cloth Band Count',
                'Procurement Axial Length [mm]', 'Fabric Area [m2]',
                'Epoxy Mass [kg]',
            )
        )
        assert actual == pytest.approx(expected)
        count_500, count_300, procurement, area, epoxy = actual
        assert procurement == 500 * count_500 + 300 * count_300
        coverage = procurement - 50 * (count_500 + count_300 - 1)
        assert coverage >= main.cell(
            row, main_columns['Total Repair Length [mm]'],
        ).value
        assert area == pytest.approx(
            3 * 3.141592653589793 * 457.2 * procurement / 1_000_000
        )
        assert epoxy == pytest.approx(area * 1.2)

    assert all(
        main.cell(6, main_columns[header]).value is None
        for header in OUTPUT_HEADERS
    )
    assert main.cell(
        7, main_columns['Wall Loss [%]'],
    ).value == pytest.approx(52.7806925498426)
    assert main.cell(7, main_columns['Repair Zone Length [mm]']).value == 300
    assert all(
        main.cell(7, main_columns[header]).value is None
        for header in (
            'Required Structural Thickness [mm]', 'Installed Plies',
            'Total Repair Length [mm]', '500 mm Cloth Band Count',
            '300 mm Cloth Band Count', 'Procurement Axial Length [mm]',
            'Fabric Area [m2]', 'Epoxy Mass [kg]',
        )
    )

    assert _warning_signature(result_book) == (
        (
            'W002',
            'No Type B Formula 12 repair solution exists for the requested '
            'case; do not install without changing the design basis or repair method.',
            '7',
        ),
        (
            'W003',
            'Requested Type B life exceeds the qualified PRW110 life; inspect, '
            'revalidate, or replace at the qualified limit.',
            '7',
        ),
        (
            'W006',
            'Type B design uses the defined through-wall defect basis and Annex F '
            'impact-qualified minimum; assessor confirmation is required.',
            '7',
        ),
    )

    assert tuple(
        cost.cell(5, column).value for column in range(1, 23)
    ) == COST_SOURCE_HEADERS
    assert tuple(
        cost.cell(5, column).value for column in range(23, 27)
    ) == ('Cost', 'Price', 'Quantity', 'Total Amount')
    assert (
        cost.tables['CostRows'].ref,
        cost.tables['CostRows'].autoFilter.ref,
    ) == ('A5:Z11', 'A5:Z11')
    assert all(
        not cost[address].protection.locked
        for address in ('B3', 'E3', 'H3', 'Y6', 'Y155')
    )
    assert all(
        cost[address].protection.locked
        for address in ('W6', 'X6', 'Z6', 'W155', 'X155', 'Z155')
    )
    assert any(
        str(item.sqref) == 'Y6:Y155'
        for item in cost.data_validations.dataValidation
    )
    expected_formulas = [
        (f'Cost Calculation!{column}{row}', formula)
        for row in range(6, 12)
        for column, formula in (
            ('W', cost_formula(row)),
            ('X', price_formula(row)),
            ('Z', total_amount_formula(row)),
        )
    ]
    assert _formula_cells(result_book) == expected_formulas
    for cost_row, main_row in zip(range(6, 12), range(2, 8), strict=True):
        assert [
            cost.cell(cost_row, column).value for column in range(1, 23)
        ] == [
            main.cell(main_row, main_columns[header]).value
            for header in COST_SOURCE_HEADERS
        ]

    main_signature = _main_result_signature(result_book)
    warning_signature = _warning_signature(result_book)
    summary_signature = _summary_signature(result_book)
    cost['B3'], cost['E3'], cost['H3'] = 25.0, 8.0, 1.4
    for row, quantity in zip(
        range(6, 12), (1, 2, 0, 3, 1.5, 4), strict=True,
    ):
        cost.cell(row, 25).value = quantity
    reupload = BytesIO()
    result_book.save(reupload)
    rebuilt = process_workbook(
        reupload.getvalue(),
        processed_at=FIXED_TIME,
        source_name='reuploaded-v13.xlsx',
    )
    rebuilt_book = load_workbook(BytesIO(rebuilt.workbook_bytes), data_only=False)

    assert rebuilt.status_counts == processed.status_counts
    assert _main_result_signature(rebuilt_book) == main_signature
    assert _warning_signature(rebuilt_book) == warning_signature
    assert _summary_signature(rebuilt_book) == summary_signature
    assert _formula_cells(rebuilt_book) == expected_formulas
    rebuilt_cost = rebuilt_book['Cost Calculation']
    assert [
        rebuilt_cost[address].value for address in ('B3', 'E3', 'H3')
    ] == [25.0, 8.0, 1.4]
    assert [
        rebuilt_cost.cell(row, 25).value for row in range(6, 12)
    ] == [1, 2, 0, 3, 1.5, 4]
    assert all(
        rebuilt_cost[address].protection.locked
        for address in ('W6', 'X6', 'Z6')
    )
    assert rebuilt_cost['Y6'].protection.locked is False
